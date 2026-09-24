"""User routes — home, search, settings, machine detail, sample."""
from io import BytesIO
from flask import Blueprint, render_template, send_file

bp = Blueprint("main", __name__)


@bp.route("/")
def home():
    return render_template("home.html")


@bp.route("/search")
def search():
    return render_template("search.html")


@bp.route("/settings")
def settings():
    return render_template("settings.html")


@bp.route("/machine/<path:code>")
def machine_detail(code):
    return render_template("machine_detail.html", machine_code=code)


@bp.route("/help/format")
def format_help():
    return render_template("format_help.html")


@bp.route("/download/sample.xlsx")
def download_sample():
    """Generate + download sample Excel file."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        return {"error": "openpyxl not installed"}, 500

    wb = Workbook()
    ws = wb.active
    ws.title = "Machines"

    # Headers
    headers = ["machine_code", "machine_name", "shop_code", "description"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Sample rows
    samples = [
        ("M_SLOT_001", "Dream Castle", "A3", "Slot machine"),
        ("M_SLOT_002", "Aladdin", "A3", "Slot machine"),
        ("R_MEDAL_001", "Avengers", "A4", "Medal game"),
        ("P_CLAW_001", "Epic Mini", "A3", "Claw machine"),
        ("A_RACING_001", "Asphalt Moto", "A4", "Racing game"),
    ]
    for row_idx, row in enumerate(samples, 2):
        for col_idx, val in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    # Column widths
    widths = [20, 25, 12, 20]
    for col_idx, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col_idx)].width = w

    # Save to BytesIO
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)

    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="machine_hub_sample.xlsx",
    )


@bp.route("/health")
def health():
    return {"status": "ok", "app": "mh-user", "mode": "offline-only"}
